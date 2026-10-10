"""Llena una copia de la plantilla `04-SABSA matrix.xlsx` desde `matriz/[1-5]-*.md`.

    uv run --with openpyxl --with pyyaml python sabsa/build_xlsx.py

Cada .md es una capa: el frontmatter trae el responsable (columna H) y cada
sección `## WHAT|WHY|HOW|WHO|WHERE|WHEN — ...` aporta a su celda los bullets que
hay bajo `**En la matriz:**`. Lo global (título, empresa) sale de `caso.yaml`.

La plantilla trae los encabezados de cada capa en las filas pares (2, 4, ..., 10)
y deja la fila impar siguiente para el contenido; la columna H es el responsable.
Se escribe sobre la plantilla original para conservar sus estilos y bordes.
"""

import re
import sys
from pathlib import Path

import yaml
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font

ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = ROOT / "04-SABSA matrix.xlsx"
CASO = ROOT / "caso.yaml"
MATRIZ = ROOT / "matriz"
OUT = ROOT / "out" / "SABSA-Matriz.xlsx"

CAPAS = ["contextual", "conceptual", "logical", "physical", "component"]
COLUMNAS = {"what": "B", "why": "C", "how": "D", "who": "E", "where": "F", "when": "G"}
FILA_CONTENIDO = {capa: 3 + 2 * i for i, capa in enumerate(CAPAS)}  # 3, 5, 7, 9, 11

ANCHO = 52  # ancho de B..G en caracteres
LINEA_PT = 13  # alto aproximado de una línea de texto en puntos


def lineas_estimadas(texto: str, ancho: int) -> int:
    """Cuántas líneas ocupa `texto` en una celda de `ancho` caracteres con wrap."""
    return sum(max(1, -(-len(linea) // ancho)) for linea in texto.split("\n"))


def sustituir(caso: dict) -> dict:
    """Reemplaza {empresa}, {area} y {proyecto} en todos los textos del caso."""
    valores = {k: caso[k] for k in ("empresa", "area", "proyecto")}

    def r(x):
        if isinstance(x, str):
            for k, v in valores.items():
                x = x.replace("{%s}" % k, v)
            return x
        if isinstance(x, list):
            return [r(i) for i in x]
        if isinstance(x, dict):
            return {k: r(v) for k, v in x.items()}
        return x

    return r(caso)


def leer_capa(path: Path) -> tuple[dict, dict[str, list[str]]]:
    """Devuelve (frontmatter, {columna: [bullets]}) de un .md de capa."""
    texto = path.read_text(encoding="utf-8")
    m = re.match(r"^---\n(.*?)\n---\n", texto, re.S)
    if not m:
        sys.exit(f"{path.name}: falta el frontmatter (--- ... ---)")
    meta = yaml.safe_load(m.group(1))
    celdas: dict[str, list[str]] = {}
    col, en_matriz = None, False
    for linea in texto[m.end():].splitlines():
        h = re.match(r"^## (WHAT|WHY|HOW|WHO|WHERE|WHEN) —", linea)
        if h:
            col, en_matriz = h.group(1).lower(), False
            if col in celdas:
                sys.exit(f"{path.name}: la sección {h.group(1)} está repetida")
            celdas[col] = []
        elif linea.startswith("## "):
            col, en_matriz = None, False
        elif col and linea.startswith("**"):
            en_matriz = linea.startswith("**En la matriz:**")
        elif col and en_matriz and linea.startswith("- "):
            # el Excel no interpreta markdown: se quitan *énfasis* y `código`
            celdas[col].append(re.sub(r"[*`]", "", linea[2:]).strip())
    return meta, celdas


def leer_matriz() -> dict[str, dict]:
    """{capa: {"responsable": str, "celdas": {columna: [bullets]}}} validado."""
    capas = {}
    for path in sorted(MATRIZ.glob("[1-5]-*.md")):
        meta, celdas = leer_capa(path)
        capas[meta["capa"]] = {"responsable": meta["responsable"], "celdas": celdas}
    errores = [f"falta la capa {c}" for c in CAPAS if c not in capas]
    for capa, datos in capas.items():
        for col in COLUMNAS:
            n = len(datos["celdas"].get(col, []))
            if n == 0:
                errores.append(f"{capa}/{col}: celda vacía")
            elif not 5 <= n <= 8:
                print(f"aviso: {capa}/{col} tiene {n} bullets (se sugieren 5–8)")
    if errores:
        sys.exit("matriz incompleta:\n  " + "\n  ".join(errores))
    return capas


def main() -> None:
    caso = yaml.safe_load(CASO.read_text(encoding="utf-8"))
    caso["capas"] = leer_matriz()
    caso = sustituir(caso)
    wb = load_workbook(TEMPLATE)
    ws = wb.active
    ws.title = "Matriz SABSA"

    # Título arriba de la matriz: se corre todo una fila hacia abajo.
    ws.insert_rows(1)
    for rango in list(ws.merged_cells.ranges):
        rango.shift(0, 1)
    ws.merge_cells("A1:H1")
    ws["A1"] = caso["titulo_matriz"]
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30
    # insert_rows no mueve las alturas de fila; se reasignan desde la plantilla.
    original = load_workbook(TEMPLATE).active
    for r in range(1, original.max_row + 1):
        ws.row_dimensions[r + 1].height = original.row_dimensions[r].height

    for col in COLUMNAS.values():
        ws.column_dimensions[col].width = ANCHO
    ws.column_dimensions["H"].width = 30

    alineacion = Alignment(wrap_text=True, vertical="top")
    for capa in CAPAS:
        fila = FILA_CONTENIDO[capa] + 1  # +1 por la fila del título
        celdas = caso["capas"][capa]["celdas"]
        max_lineas = 1
        for clave, col in COLUMNAS.items():
            texto = "\n".join(f"- {item}" for item in celdas[clave])
            celda = ws[f"{col}{fila}"]
            celda.value = texto
            celda.font = Font(bold=False, size=10)
            celda.alignment = alineacion
            max_lineas = max(max_lineas, lineas_estimadas(texto, ANCHO))

        resp = ws[f"H{fila}"]
        resp.value = caso["capas"][capa]["responsable"]
        resp.font = Font(bold=True, size=10)
        resp.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.row_dimensions[fila].height = min(409, max_lineas * LINEA_PT + 8)

    ws.freeze_panes = "B3"
    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    print(f"ok → {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
